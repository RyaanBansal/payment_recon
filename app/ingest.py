"""Reading the uploaded .xlsx files and loading them into SQLite."""
import io
import warnings
from datetime import datetime, time

from openpyxl import load_workbook

from .db import conn, q, get_meta, set_meta, table_exists, now

warnings.filterwarnings("ignore", message="Workbook contains no default style")

PAY_REQUIRED = ["CustomerPayment ID", "InvoicePayment ID", "Reference Number", "Amount", "Mode", "Deposit To", "Date"]
ZOHO_REQUIRED = ["id", "amount", "notes", "method", "created_at"]
ZOHO_FMT = "%d/%m/%Y %H:%M:%S"  # the display format used in the Razorpay/Zoho sheet


class SheetError(ValueError):
    pass


# ---------- reading ----------
def read_raw(data: bytes):
    """Return (headers, rows) with raw cell values. Duplicate headers get a .1/.2 suffix (as pandas does)."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Could not open the file as an .xlsx workbook ({e}).")
    it = wb.worksheets[0].iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        raise SheetError("The sheet is empty.")
    headers, seen = [], {}
    for h in header:
        h = "" if h is None else str(h).strip()
        if h in seen:
            seen[h] += 1
            h = f"{h}.{seen[h]}"
        else:
            seen[h] = 0
        headers.append(h)
    n = len(headers)
    rows = [(list(r) + [None] * n)[:n] for r in it if any(v not in (None, "") for v in r)]  # pad short rows
    return headers, rows


def to_text(v):
    """Cell -> text, keeping the format the sheet displays (IDs never go through float)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).strip()
    return s or None


def to_float(s):
    try:
        return float(str(s).replace(",", ""))
    except (TypeError, ValueError):
        return None


def normalize_amount(amount, currency=None, exchange_rate=None):
    """Return the amount converted using the sheet's exchange rate when present."""
    try:
        value = float(str(amount).replace(",", "")) if amount not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0

    if exchange_rate in (None, "", 0):
        return value

    try:
        rate = float(str(exchange_rate).replace(",", ""))
    except (TypeError, ValueError):
        return value

    if rate == 0:
        return value
    return value * rate


def parse_dt(v):
    if isinstance(v, datetime):
        return v
    if not isinstance(v, str):
        return None
    for f in (ZOHO_FMT, "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(v.strip(), f)
        except ValueError:
            pass
    return None


def _check(headers, required, label):
    missing = [c for c in required if c not in headers]
    if missing:
        raise SheetError(f"This doesn't look like the {label} sheet - missing column(s): {', '.join(missing)}.")


def _ensure_columns(c, table, cols):
    have = {r["name"] for r in c.execute(f"PRAGMA table_info({q(table)})")}
    for col in cols:
        if col not in have:
            c.execute(f"ALTER TABLE {q(table)} ADD COLUMN {q(col)}")


# ---------- payments (customer payment sheet: full list -> replaces previous data) ----------
def load_payments(files) -> dict:
    """Load one or more Zoho (customer payment) sheets in one go.

    `files` is a list of (filename, bytes). The previous Zoho data is replaced by the combined contents of the
    files. Every file is checked before the database is touched, so a bad file leaves the old data intact.
    Rows in a later file that are identical to a row from an earlier file (overlapping exports) are skipped.
    A single `bytes` value is accepted too.
    """
    if isinstance(files, (bytes, bytearray)):
        files = [("Zoho sheet", bytes(files))]

    parsed = []
    for name, data in files:
        try:
            headers, rows = read_raw(data)
            _check(headers, PAY_REQUIRED, "Zoho")
        except SheetError as e:
            raise SheetError(f"{name}: {e}")
        parsed.append((name, headers, rows))

    cols = []  # union of the files' columns, in first-seen order
    for _, headers, _ in parsed:
        for h in headers:
            if h and h not in cols:
                cols.append(h)
    derived = ["_date", "_amount"]
    out, seen, per_file, dupes = [], set(), [], 0
    for name, headers, rows in parsed:
        idx = {h: headers.index(h) for h in headers if h}
        added = 0
        file_seen = set()
        for r in rows:
            vals = [to_text(r[idx[h]]) if h in idx else None for h in cols]
            key = tuple(vals)
            if key in seen:  # already loaded from an earlier file in this batch
                dupes += 1
                continue
            file_seen.add(key)
            d = parse_dt(r[idx["Date"]])
            raw_amount = to_float(to_text(r[idx["Amount"]]))
            currency = to_text(r[idx["Currency"]]) if "Currency" in idx else None
            exchange_rate = to_float(to_text(r[idx["Exchange Rate"]])) if "Exchange Rate" in idx else None
            out.append(vals + [d.strftime("%Y-%m-%d") if d else None, normalize_amount(raw_amount, currency, exchange_rate)])
            added += 1
        seen |= file_seen
        per_file.append((name, added))

    with conn() as c:
        c.execute("DROP TABLE IF EXISTS payments")
        c.execute(f"CREATE TABLE payments ({', '.join(q(x) for x in cols + derived)})")
        for n, name in enumerate(["Reference Number", "InvoicePayment ID", "CustomerPayment ID", "Mode", "Deposit To", "_date"]):
            c.execute(f"CREATE INDEX ix_pay_{n} ON payments({q(name)})")
        c.executemany(f"INSERT INTO payments VALUES ({','.join('?' * (len(cols) + len(derived)))})", out)
        set_meta(c, "payments_cols", cols)
        set_meta(c, "data_updated_at", now())
        set_meta(c, "payments_loaded_at", now())
    warns = [f"{dupes:,} duplicate row(s) that appeared in more than one file were loaded only once."] if dupes else []
    n_files = len(parsed)
    msg = (f"Loaded {len(out):,} rows from the Zoho sheet (previous Zoho sheet data replaced)." if n_files == 1 else
           f"Loaded {len(out):,} rows from {n_files} Zoho sheets (previous Zoho sheet data replaced).")
    return {"rows": len(out), "message": msg, "warnings": warns,
            "results": [{"filename": n, "ok": True, "message": f"{k:,} rows"} for n, k in per_file]}


# ---------- zoho / razorpay (monthly sheets -> accumulate, upsert on id) ----------
def load_zoho(data: bytes) -> dict:
    headers, rows = read_raw(data)
    _check(headers, ZOHO_REQUIRED, "Razorpay")
    keep = [i for i, h in enumerate(headers) if h]
    cols = [headers[i] for i in keep]
    idx = {h: headers.index(h) for h in cols}

    # created_at: text is dd/mm/yyyy, but Excel turned some cells into real datetimes with day and month
    # swapped (it read the text in the wrong locale). If a sheet mixes both kinds, swap those back.
    vals = [r[idx["created_at"]] for r in rows]
    n_txt = sum(isinstance(v, str) for v in vals)
    n_dt = sum(isinstance(v, datetime) for v in vals)
    swap = n_txt > 0 and n_dt > 0
    swapped = 0
    created_disp, created_iso = [], []
    for v in vals:
        if isinstance(v, datetime) and swap and v.day <= 12:
            v = v.replace(month=v.day, day=v.month)
            swapped += 1
        d = parse_dt(v)
        created_disp.append(d.strftime(ZOHO_FMT) if d else to_text(v))
        created_iso.append(d.strftime("%Y-%m-%d %H:%M:%S") if d else None)
    bad_dates = sum(1 for v, d in zip(vals, created_iso) if v not in (None, "") and d is None)

    derived = ["_created", "_amount", "_fee"]
    all_cols = cols + derived
    with conn() as c:
        if not table_exists(c, "zoho"):
            c.execute('CREATE TABLE zoho ("id" TEXT PRIMARY KEY)')
            c.execute('CREATE INDEX ix_zoho_created ON zoho("_created")')
        _ensure_columns(c, "zoho", all_cols)
        existing = {r["id"] for r in c.execute("SELECT id FROM zoho")}
        out, ids = [], set()
        for n, r in enumerate(rows):
            row = [to_text(r[i]) for i in keep]
            row[cols.index("created_at")] = created_disp[n]
            pid = row[cols.index("id")]
            if not pid:
                continue
            ids.add(pid)
            fee = to_float(to_text(r[idx["fee"]])) if "fee" in idx else None
            raw_amount = to_float(to_text(r[idx["amount"]]))
            currency = to_text(r[idx["Currency"]]) if "Currency" in idx else None
            exchange_rate = to_float(to_text(r[idx["Exchange Rate"]])) if "Exchange Rate" in idx else None
            out.append(row + [created_iso[n], normalize_amount(raw_amount, currency, exchange_rate), fee])
        c.executemany(
            f"INSERT OR REPLACE INTO zoho ({', '.join(q(x) for x in all_cols)}) VALUES ({','.join('?' * len(all_cols))})",
            out,
        )
        prev = get_meta(c, "zoho_cols", [])
        set_meta(c, "zoho_cols", prev + [x for x in cols if x not in prev])
        set_meta(c, "data_updated_at", now())
        set_meta(c, "zoho_loaded_at", now())
    updated = len(ids & existing)
    warns = []
    if swapped:
        warns.append(f"{swapped:,} created_at cells were real dates with day/month swapped by Excel; corrected to dd/mm/yyyy.")
    if bad_dates:
        warns.append(f"{bad_dates:,} created_at values could not be read as dates (excluded from date filters).")
    return {"rows": len(ids), "message": f"Loaded {len(ids):,} rows from the Razorpay sheet ({len(ids) - updated:,} new, {updated:,} updated).", "warnings": warns}