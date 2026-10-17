"""Loading the bank statements (Kotak, ICICI, HDFC) and the Cash Free transactions sheet into SQLite.

All are monthly / period sheets: uploads accumulate and re-uploading a period is safe (rows are upserted on a key
built from the row's own content, `_key`). The sheet's own columns are stored as they are (dates are shown as
dd-mm-yyyy); the extra `_` columns are derived: `_dt` (ISO date-time), `_amount`, `_ref` (the matching id,
normalised) and a few sheet specific ones (`_drcr` = CR / DR for the bank statements).
"""
import csv
import hashlib
import io
import re
import warnings
from datetime import datetime

from openpyxl import load_workbook

from .db import conn, get_meta, now, q, set_meta, table_exists
from .ingest import SheetError, _ensure_columns, to_float, to_text
from .match_common import norm_ref

warnings.filterwarnings("ignore", message="Workbook contains no default style")

DISPLAY_FMT = "%d-%m-%Y %H:%M"  # how the sheets show their date-time
DISPLAY_DATE = "%d-%m-%Y"       # ... and their date


# ---------- reading (csv or xlsx; xlsx -> every sheet) ----------
def canon(header) -> str:
    """Header -> comparison key: 'Chq /Ref No.' -> 'chqrefno', 'Phone NO' and 'Phone No.' -> 'phoneno'."""
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


def _headers(header_row):
    """(display names, canonical keys). Blank headers stay blank; repeated headers ('Dr / Cr' twice) get
    a .1 / .2 suffix on both, so each column keeps its own identity."""
    shown, keys, seen = [], [], {}
    for h in header_row:
        h = "" if h is None else " ".join(str(h).split())  # also removes stray tabs / double spaces
        k = canon(h)
        if h and k in seen:
            seen[k] += 1
            h, k = f"{h}.{seen[k]}", f"{k}.{seen[k]}"
        elif h:
            seen[k] = 0
        shown.append(h)
        keys.append(k)
    return shown, keys


def _pad(rows, n):
    return [(list(r) + [None] * n)[:n] for r in rows if any(v not in (None, "") for v in r)]


def _decode_csv(data: bytes):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise SheetError("Could not read the file as text.")


def _xlsx_rows(ws):
    # some banks write a wrong sheet-size record (A1), which makes a read-only workbook look empty
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _xls_grids(data: bytes):
    """Old .xls workbook (HDFC) -> [(sheet name, rows)]."""
    try:
        import xlrd
    except ImportError:
        raise SheetError("Reading .xls files needs the 'xlrd' package - run: pip install xlrd")
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Could not open the file as an .xls workbook ({e}).")
    out = []
    for sh in book.sheets():
        rows = []
        for i in range(sh.nrows):
            row = []
            for cell in sh.row(i):
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def read_grids(data: bytes, filename: str = ""):
    """-> [(sheet name, rows)] with every row exactly as in the file (no header assumed). .csv / .xlsx / .xls."""
    name = filename.lower()
    if name.endswith(".csv"):
        return [("csv", [r for r in csv.reader(io.StringIO(_decode_csv(data)))])]
    if name.endswith(".xls"):
        return _xls_grids(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Could not open the file as an .xlsx workbook ({e}).")
    return [(ws.title, _xlsx_rows(ws)) for ws in wb.worksheets]


def read_statement(data: bytes, filename: str, required: dict, label: str):
    """Bank statements have the account details above the table and notes below it. Find the row that holds
    the column headers (all `required` columns present) and return (display headers, canonical keys, rows
    below it - blank lines dropped; notes / totals are left for the caller to skip)."""
    want = set(required.values())
    for _, rows in read_grids(data, filename):
        for i, row in enumerate(rows):
            shown, keys = _headers(row)
            if want <= set(keys):
                return shown, keys, _pad(rows[i + 1:], len(shown))
    raise SheetError(f"This doesn't look like the {label} statement - could not find a header row with: "
                     f"{', '.join(required)}.")


def read_sheets(data: bytes, filename: str = ""):
    """-> [(sheet name, display headers, canonical keys, rows)]. A .csv gives one entry."""
    if filename.lower().endswith(".csv"):
        text = _decode_csv(data)
        rows = [r for r in csv.reader(io.StringIO(text)) if any(v.strip() for v in r)]
        if not rows:
            raise SheetError("The file is empty.")
        shown, keys = _headers(rows[0])
        return [("csv", shown, keys, _pad(rows[1:], len(shown)))]
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SheetError(f"Could not open the file as an .xlsx workbook ({e}).")
    out = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if header is None or not any(header):
            continue
        shown, keys = _headers(header)
        out.append((ws.title, shown, keys, _pad(it, len(shown))))
    if not out:
        raise SheetError("The workbook is empty.")
    return out


def _missing(keys, required):
    have = set(keys)
    return [name for name, k in required.items() if k not in have]


DATE_FORMATS = ("%d-%m-%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
                "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                "%d/%m/%Y %I:%M:%S %p",  # ICICI posted date: 02/04/2026 08:31:55 PM
                "%d/%b/%Y",              # ICICI: 02/Apr/2026
                "%d/%m/%y")              # HDFC: 02/04/26


def parse_when(v):
    """Cell -> datetime (text in the sheets' dd-mm-yyyy style, or a real Excel date)."""
    if isinstance(v, datetime):
        return v
    if not isinstance(v, str):
        return None
    for f in DATE_FORMATS:
        try:
            return datetime.strptime(v.strip(), f)
        except ValueError:
            pass
    return None


def _sha(*parts) -> str:
    return hashlib.sha1("|".join("" if p is None else str(p) for p in parts).encode()).hexdigest()[:20]


def _store(table, cols, derived, rows, meta_prefix):
    """Upsert rows (each: [_key] + values of cols + values of derived) into `table`. Returns (new, updated)."""
    all_cols = cols + derived
    with conn() as c:
        if not table_exists(c, table):
            c.execute(f'CREATE TABLE {q(table)} ("_key" TEXT PRIMARY KEY)')
            c.execute(f'CREATE INDEX {q("ix_" + table + "_dt")} ON {q(table)}("_dt")')
        _ensure_columns(c, table, all_cols)
        existing = {r[0] for r in c.execute(f'SELECT "_key" FROM {q(table)}')}
        c.executemany(
            f'INSERT OR REPLACE INTO {q(table)} ("_key", {", ".join(q(x) for x in all_cols)}) VALUES ({",".join("?" * (len(all_cols) + 1))})',
            rows)
        prev = get_meta(c, f"{meta_prefix}_cols", [])
        set_meta(c, f"{meta_prefix}_cols", prev + [x for x in cols if x not in prev])
        set_meta(c, "data_updated_at", now())
        set_meta(c, f"{meta_prefix}_loaded_at", now())
    keys = {r[0] for r in rows}
    updated = len(keys & existing)
    return len(keys) - updated, updated


def _iso(when):
    return when.strftime("%Y-%m-%d %H:%M:%S") if when else None


# ---------- Kotak ----------
KOTAK_REQUIRED = {"Transaction Date": "transactiondate", "Chq /Ref No.": "chqrefno", "Amount": "amount", "Dr / Cr": "drcr"}


def _kotak_core(shown, keys, rows, table, label) -> dict:
    """Shared Kotak parsing for the downloaded monthly sheet data."""
    miss = _missing(keys, KOTAK_REQUIRED)
    if miss:
        raise SheetError(f"This doesn't look like the {label} statement - missing column(s): {', '.join(miss)}.")
    keep = [i for i, h in enumerate(shown) if h]
    cols = [shown[i] for i in keep]
    at = {k: keys.index(k) for k in keys if k}  # canonical key -> column index (first one)
    i_date, i_ref, i_amt, i_dc = at["transactiondate"], at["chqrefno"], at["amount"], at["drcr"]
    i_val, i_desc, i_bal = at.get("valuedate"), at.get("description"), at.get("balance")

    out, skipped, dupes, seen = [], 0, 0, set()
    for r in rows:
        amount = to_float(to_text(r[i_amt]))
        if amount is None:  # 'Closing Balance ...' footer or any other line that is not a transaction
            skipped += 1
            continue
        vals = [to_text(r[i]) for i in keep]
        when = parse_when(r[i_date] if isinstance(r[i_date], datetime) else to_text(r[i_date])) \
            or (parse_when(r[i_val] if isinstance(r[i_val], datetime) else to_text(r[i_val])) if i_val is not None else None)
        for i in (i_date, i_val):  # real Excel dates -> the sheet's dd-mm-yyyy display
            if i is not None and isinstance(r[i], datetime):
                vals[keep.index(i)] = r[i].strftime(DISPLAY_FMT)
        ref = to_text(r[i_ref])
        key = _sha(*(to_text(r[i]) for i in (i_date, i_val, i_desc, i_ref, i_amt, i_dc, i_bal) if i is not None))
        if key in seen:
            dupes += 1
        seen.add(key)
        out.append([key] + vals + [_iso(when), amount, norm_ref(ref), (to_text(r[i_dc]) or "").upper()])
    if not out:
        raise SheetError("No transactions found in the statement.")
    new, updated = _store(table, cols, ["_dt", "_amount", "_ref", "_drcr"], out, table)
    warns = []
    if skipped:
        warns.append(f"{skipped:,} line(s) that are not transactions (e.g. the closing balance and notes) were skipped.")
    if dupes:
        warns.append(f"{dupes:,} identical row(s) were loaded only once.")
    return {"rows": len(seen), "message": f"Loaded {len(seen):,} rows from the {label} sheet ({new:,} new, {updated:,} updated).",
            "warnings": warns}


def load_kotak(data: bytes, filename: str = "") -> dict:
    _, shown, keys, rows = read_sheets(data, filename)[0]
    return _kotak_core(shown, keys, rows, "kotak", "Kotak")


# ---------- ICICI ----------
ICICI_REQUIRED = {"Tran. Id": "tranid", "Transaction Date": "transactiondate",
                  "Withdrawal Amt (INR)": "withdrawalamtinr", "Deposit Amt (INR)": "depositamtinr"}


def _display_dates(vals, keep, r, indexes, fmt_for):
    """Rewrite the date cells of one row to dd-mm-yyyy (keeps what the other sheets show)."""
    for i in indexes:
        if i is None:
            continue
        d = parse_when(r[i] if isinstance(r[i], datetime) else to_text(r[i]))
        if d:
            vals[keep.index(i)] = d.strftime(fmt_for(i))


def load_icici(data: bytes, filename: str = "") -> dict:
    shown, keys, rows = read_statement(data, filename, ICICI_REQUIRED, "ICICI")
    keep = [i for i, h in enumerate(shown) if h]
    cols = [shown[i] for i in keep]
    at = {k: keys.index(k) for k in keys if k}
    i_id, i_tdate, i_dep, i_wd = at["tranid"], at["transactiondate"], at["depositamtinr"], at["withdrawalamtinr"]
    i_vdate, i_posted, i_bal = at.get("valuedate"), at.get("transactionposteddate"), at.get("balanceinr")
    i_sn = at.get("sn")

    out, seen, no_amount, dupes = [], set(), 0, 0
    for r in rows:
        tran_id = to_text(r[i_id])
        if not tran_id or (i_sn is not None and not (to_text(r[i_sn]) or "").isdigit()):
            continue  # the summary (opening / closing balance) and the legend of transaction codes at the bottom
        dep, wd = to_float(to_text(r[i_dep])), to_float(to_text(r[i_wd]))
        if dep is None and wd is None:
            no_amount += 1
            continue
        credit = dep is not None
        amount = dep if credit else wd
        vals = [to_text(r[i]) for i in keep]
        when = (parse_when(r[i_posted] if isinstance(r[i_posted], datetime) else to_text(r[i_posted])) if i_posted is not None else None) \
            or parse_when(r[i_tdate] if isinstance(r[i_tdate], datetime) else to_text(r[i_tdate]))
        _display_dates(vals, keep, r, (i_vdate, i_tdate, i_posted), lambda i: DISPLAY_FMT if i == i_posted else DISPLAY_DATE)
        key = _sha(tran_id, to_text(r[i_dep]), to_text(r[i_wd]), to_text(r[i_bal]) if i_bal is not None else None)
        if key in seen:
            dupes += 1
        seen.add(key)
        out.append([key] + vals + [_iso(when), amount, norm_ref(tran_id), "CR" if credit else "DR"])
    if not out:
        raise SheetError("No transactions found in the statement.")
    new, updated = _store("icici", cols, ["_dt", "_amount", "_ref", "_drcr"], out, "icici")
    warns = []
    if no_amount:
        warns.append(f"{no_amount:,} line(s) without a deposit or withdrawal amount were skipped.")
    if dupes:
        warns.append(f"{dupes:,} identical row(s) were loaded only once.")
    return {"rows": len(seen), "message": f"Loaded {len(seen):,} rows from the ICICI sheet ({new:,} new, {updated:,} updated).",
            "warnings": warns}


# ---------- HDFC ----------
HDFC_REQUIRED = {"Date": "date", "Narration": "narration", "Chq./Ref.No.": "chqrefno", "Value Dt": "valuedt",
                 "Withdrawal Amt.": "withdrawalamt", "Deposit Amt.": "depositamt"}


def load_hdfc(data: bytes, filename: str = "") -> dict:
    shown, keys, rows = read_statement(data, filename, HDFC_REQUIRED, "HDFC")
    keep = [i for i, h in enumerate(shown) if h]
    cols = [shown[i] for i in keep]
    at = {k: keys.index(k) for k in keys if k}
    i_date, i_ref, i_dep, i_wd = at["date"], at["chqrefno"], at["depositamt"], at["withdrawalamt"]
    i_val, i_nar, i_bal = at["valuedt"], at["narration"], at.get("closingbalance")

    out, seen, no_amount, dupes = [], set(), 0, 0
    for r in rows:
        when = parse_when(r[i_date] if isinstance(r[i_date], datetime) else to_text(r[i_date]))
        if when is None:  # '****' separators, the statement summary and the notes at the bottom
            continue
        dep, wd = to_float(to_text(r[i_dep])), to_float(to_text(r[i_wd]))
        if dep is None and wd is None:
            no_amount += 1
            continue
        credit = dep is not None
        vals = [to_text(r[i]) for i in keep]
        _display_dates(vals, keep, r, (i_date, i_val), lambda i: DISPLAY_DATE)
        ref = to_text(r[i_ref])
        # '0000581598905825' -> '581598905825' (the UPI reference); the sheet's own text stays in the table
        norm = re.sub(r"^0+", "", norm_ref(ref))
        key = _sha(to_text(r[i_date]), to_text(r[i_val]), to_text(r[i_nar]), ref, to_text(r[i_dep]), to_text(r[i_wd]),
                   to_text(r[i_bal]) if i_bal is not None else None)
        if key in seen:
            dupes += 1
        seen.add(key)
        out.append([key] + vals + [_iso(when), dep if credit else wd, norm, "CR" if credit else "DR"])
    if not out:
        raise SheetError("No transactions found in the statement.")
    new, updated = _store("hdfc", cols, ["_dt", "_amount", "_ref", "_drcr"], out, "hdfc")
    warns = []
    if no_amount:
        warns.append(f"{no_amount:,} line(s) without a deposit or withdrawal amount were skipped.")
    if dupes:
        warns.append(f"{dupes:,} identical row(s) were loaded only once.")
    return {"rows": len(seen), "message": f"Loaded {len(seen):,} rows from the HDFC sheet ({new:,} new, {updated:,} updated).",
            "warnings": warns}


# ---------- Cash Free ----------
CASHFREE_REQUIRED = {"Transaction Time": "transactiontime", "Order Id": "orderid", "Order Amount": "orderamount"}


def load_cashfree(data: bytes, filename: str = "") -> dict:
    sheets, skipped_sheets = [], []
    for name, shown, keys, rows in read_sheets(data, filename):
        (skipped_sheets if _missing(keys, CASHFREE_REQUIRED) else sheets).append((name, shown, keys, rows))
    if not sheets:
        name, _, keys, _ = skipped_sheets[0]
        raise SheetError("This doesn't look like the Cash Free sheet - missing column(s): "
                         f"{', '.join(_missing(keys, CASHFREE_REQUIRED))}.")

    # the month tabs don't all have the same columns (and spell some differently): union them by canonical name
    cols, ckeys = [], []
    for _, shown, keys, _ in sheets:
        for h, k in zip(shown, keys):
            if h and k not in ckeys:
                cols.append(h)
                ckeys.append(k)

    # Transaction Time: text is dd-mm-yyyy, but Excel turned some cells into real dates with day and month swapped
    # (it read the text in the wrong locale). If the workbook has both kinds, those dates are swapped back.
    all_t = [r[keys.index("transactiontime")] for _, _, keys, rows in sheets for r in rows]
    swap = any(isinstance(v, str) for v in all_t) and any(isinstance(v, datetime) for v in all_t)

    out, seen, swapped, no_id, bad_dates = [], set(), 0, 0, 0
    for _, shown, keys, rows in sheets:
        at = {k: keys.index(k) for k in keys if k}
        for r in rows:
            vals = {k: to_text(r[i]) for k, i in at.items()}
            order_id = vals.get("orderid")
            if not order_id:
                no_id += 1
                continue
            v = r[at["transactiontime"]]
            if isinstance(v, datetime) and swap and v.day <= 12:
                v = v.replace(month=v.day, day=v.month)
                swapped += 1
            when = parse_when(v)
            if when:
                vals["transactiontime"] = when.strftime(DISPLAY_FMT)
            elif vals.get("transactiontime"):
                bad_dates += 1
            key = _sha(order_id, vals.get("cashfreeorderid"), vals.get("referenceid"), vals.get("transactionstatus"),
                       vals.get("transactionamount"))
            seen.add(key)
            order_amt = to_float(vals.get("orderamount"))
            txn_amt = to_float(vals.get("transactionamount"))
            out.append([key] + [vals.get(k) for k in ckeys] + [_iso(when),
                                                                  order_amt if order_amt is not None else 0.0, txn_amt,
                                                                  norm_ref(order_id)])
    if not out:
        raise SheetError("No transactions found in the sheet.")
    new, updated = _store("cashfree", cols, ["_dt", "_amount", "_txn_amount", "_ref"], out, "cashfree")
    warns = []
    if len(sheets) > 1:
        warns.append(f"Read {len(sheets)} month tabs: {', '.join(s[0] for s in sheets)}.")
    if skipped_sheets:
        warns.append(f"Skipped tab(s) without the Cash Free columns: {', '.join(s[0] for s in skipped_sheets)}.")
    if swapped:
        warns.append(f"{swapped:,} Transaction Time cells were real dates with day/month swapped by Excel; corrected to dd-mm-yyyy.")
    if no_id:
        warns.append(f"{no_id:,} row(s) without an Order Id were skipped.")
    if bad_dates:
        warns.append(f"{bad_dates:,} Transaction Time values could not be read as dates (excluded from date filters).")
    return {"rows": len(seen), "message": f"Loaded {len(seen):,} rows from the Cash Free sheet ({new:,} new, {updated:,} updated).",
            "warnings": warns}